"""Excel Quick View: an in-app window that shows only the matched rows.

The search results table reports spreadsheet matches as ``Sheet!A1`` locations.
Double-clicking such a row offers two choices (Quick View / open in Excel); this
module implements the Quick View side.

Only the sheets that contain a match are shown, and inside a sheet only the
matched rows are listed. The first row of every sheet is treated as the header
and becomes the table's column titles (Excel's A/B/C letters are *not* shown).
The real Excel row numbers are kept in a leading ``#`` column so the user can
find the row in Excel again.
"""

import os

from PySide6.QtCore import Qt
from PySide6.QtGui import QBrush, QColor
from PySide6.QtWidgets import (
    QApplication,
    QDialog,
    QHBoxLayout,
    QLabel,
    QPushButton,
    QTabWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from .matching import matches

# Spreadsheet extensions handled by the Open XML reader (openpyxl) and the
# legacy binary reader (xlrd). Mirrors ``extractors._TYPE_TO_EXTRACTOR``.
OPENXML_EXTENSIONS = {".xlsx", ".xlsm", ".xltx"}
LEGACY_EXTENSIONS = {".xls"}
EXCEL_EXTENSIONS = OPENXML_EXTENSIONS | LEGACY_EXTENSIONS

# How many matched rows are listed per sheet. Beyond this the footer reports
# that the view was truncated, keeping the dialog responsive on huge books.
_MAX_ROWS = 500

# Highlight applied to the cells that actually contain the match.
_HIT_BRUSH = QBrush(QColor("#fff2b2"))


def is_excel_path(path: str) -> bool:
    """True when ``path`` points at a spreadsheet Quick View can read."""
    return os.path.splitext(path)[1].lower() in EXCEL_EXTENSIONS


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _column_letter(index: int) -> str:
    """Excel-style column letter for a zero-based column index."""
    letters = ""
    index += 1
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def _header_titles(header_values):
    """Column titles from Excel's first row.

    Empty header cells fall back to their column letter (A, B, C, …) so the
    table never shows a blank title. The letters are only used as a fallback,
    never as the primary label.
    """
    titles = []
    for index, value in enumerate(header_values):
        title = _cell_text(value)
        titles.append(title or _column_letter(index))
    return titles


def _sheet_rows(rows, query, method, threshold):
    """Reduce a sheet to its matched rows.

    ``rows`` is an iterable of raw cell values (the header row included). The
    first row is returned as the header, and every following row that contains
    at least one matching cell is kept along with the indices of the hits.
    """
    rows = list(rows)
    if not rows:
        return None

    headers = _header_titles(rows[0])
    width = max(len(headers), max((len(r) for r in rows), default=0))
    headers += [_column_letter(i) for i in range(len(headers), width)]

    matched = []
    for offset, raw in enumerate(rows[1:], start=2):
        values = [_cell_text(v) for v in raw]
        values += [""] * (width - len(values))
        hits = [i for i, text in enumerate(values)
                if text and matches(query, text, method, threshold)]
        if hits:
            matched.append({"row": offset, "values": values, "hits": hits})
            if len(matched) >= _MAX_ROWS:
                break

    if not matched:
        return None
    return {"headers": headers, "rows": matched}


def _openpyxl_sheets(path):
    """Yield ``(name, rows)`` for every sheet of an Open XML workbook."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            yield sheet.title, sheet.iter_rows(values_only=True)
    finally:
        workbook.close()


def _xlrd_sheets(path):
    """Yield ``(name, rows)`` for every sheet of a legacy .xls workbook."""
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        for sheet in workbook.sheets():
            rows = (sheet.row_values(index) for index in range(sheet.nrows))
            yield sheet.name, rows
    finally:
        workbook.release_resources()


def load_matches(path, query, method, threshold):
    """Return the matched rows of an Excel file, grouped by sheet.

    The result maps a sheet name to ``{"headers": [...], "rows": [...]}`` where
    each row is ``{"row": <excel row number>, "values": [...], "hits": [...]}``.
    Only sheets that contain at least one match are included.

    Raises whatever the underlying reader raises (missing library, corrupt
    file); the caller falls back to opening the file externally.
    """
    extension = os.path.splitext(path)[1].lower()
    if extension in OPENXML_EXTENSIONS:
        source = _openpyxl_sheets(path)
    elif extension in LEGACY_EXTENSIONS:
        source = _xlrd_sheets(path)
    else:
        return {}

    result = {}
    for name, rows in source:
        sheet = _sheet_rows(rows, query, method, threshold)
        if sheet:
            result[name] = sheet
    return result



class _SheetTable(QWidget):
    """Read-only table listing the matched rows of one sheet."""

    def __init__(self, sheet, i18n, parent=None):
        super().__init__(parent)
        self._i18n = i18n

        box = QVBoxLayout(self)
        box.setContentsMargins(6, 6, 6, 6)
        box.setSpacing(4)

        headers = sheet["headers"]
        rows = sheet["rows"]

        self.table = QTableWidget(len(rows), len(headers) + 1)
        self.table.setHorizontalHeaderLabels(
            [self._i18n.t("excel_view_row_header")] + headers)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows)
        self.table.setAlternatingRowColors(True)
        self.table.setWordWrap(False)
        self.table.verticalHeader().setVisible(False)

        for row_index, entry in enumerate(rows):
            number_item = QTableWidgetItem(str(entry["row"]))
            number_item.setFlags(Qt.ItemFlag.ItemIsEnabled)
            number_item.setForeground(QBrush(QColor("#777777")))
            self.table.setItem(row_index, 0, number_item)
            for col_index, value in enumerate(entry["values"], start=1):
                item = QTableWidgetItem(value)
                item.setFlags(Qt.ItemFlag.ItemIsEnabled)
                if col_index - 1 in entry["hits"]:
                    item.setBackground(_HIT_BRUSH)
                    font = item.font()
                    font.setBold(True)
                    item.setFont(font)
                self.table.setItem(row_index, col_index, item)

        self.table.resizeColumnsToContents()
        # Keep the leading row-number column narrow and stable.
        self.table.setColumnWidth(0, max(48, self.table.columnWidth(0)))
        box.addWidget(self.table, 1)

        self.status = QLabel()
        self.status.setStyleSheet("color:#555;")
        box.addWidget(self.status)
        self._apply_status(sheet)

    def _apply_status(self, sheet):
        text = self._i18n.t("excel_view_matches", n=len(sheet["rows"]))
        if len(sheet["rows"]) >= _MAX_ROWS:
            text += "  " + self._i18n.t("excel_view_truncated", n=_MAX_ROWS)
        self.status.setText(text)


class ExcelQuickView(QDialog):
    """Modal window showing the matched Excel rows, grouped by sheet."""

    def __init__(self, path, query, method, threshold, i18n, parent=None):
        super().__init__(parent)
        self._path = path
        self._i18n = i18n
        self._sheets = load_matches(path, query, method, threshold)

        self.setWindowTitle(
            self._i18n.t("excel_view_title", name=os.path.basename(path)))
        self.setModal(True)

        box = QVBoxLayout(self)
        box.setContentsMargins(8, 8, 8, 8)
        box.setSpacing(6)

        if self._sheets:
            self.tabs = QTabWidget()
            # Sheet tabs sit at the bottom-left, mirroring Excel.
            self.tabs.setTabPosition(QTabWidget.TabPosition.South)
            for name, sheet in self._sheets.items():
                self.tabs.addTab(_SheetTable(sheet, i18n), name)
            box.addWidget(self.tabs, 1)
        else:
            self.tabs = None
            empty = QLabel(self._i18n.t("excel_view_no_matches"))
            empty.setAlignment(Qt.AlignmentFlag.AlignCenter)
            box.addWidget(empty, 1)

        buttons = QHBoxLayout()
        buttons.addStretch()
        self.button_close = QPushButton(self._i18n.t("close_button"))
        self.button_close.clicked.connect(self.accept)
        buttons.addWidget(self.button_close)
        box.addLayout(buttons)

        self._apply_geometry()

    def _apply_geometry(self):
        """Fill 90% of the screen, centred, never spilling off the display."""
        screen = self.screen() or QApplication.primaryScreen()
        if screen is None:
            self.resize(900, 600)
            return
        available = screen.availableGeometry()
        width = int(available.width() * 0.9)
        height = int(available.height() * 0.9)
        self.setMinimumSize(360, 240)
        self.resize(width, height)
        self.move(
            available.x() + (available.width() - width) // 2,
            available.y() + (available.height() - height) // 2,
        )

    def showEvent(self, event):
        # Re-centre whenever the dialog is shown, so it lands on the display
        # the user is actually working on.
        self._apply_geometry()
        super().showEvent(event)

    @property
    def sheet_names(self):
        """Names of the sheets currently shown (only those with matches)."""
        if self.tabs is None:
            return []
        return [self.tabs.tabText(i) for i in range(self.tabs.count())]
