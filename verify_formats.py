"""Ad-hoc verification for spreadsheet/Word format coverage.

Not part of the shipped application: run manually with
``python verify_formats.py``.
"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from src import extractors  # noqa: E402
from src.i18n import I18n  # noqa: E402
from src.scanner import ScanWorker  # noqa: E402

TYPES = {"filename": True, "word": True, "excel": True, "pdf": True,
         "txt": True}


def test_every_accepted_extension_has_an_extractor():
    """No format may be accepted by the UI yet silently skipped."""
    accepted = set()
    for extensions in extractors.TYPE_TO_EXTENSIONS.values():
        accepted |= extensions
    missing = sorted(ext for ext in accepted
                     if ext not in extractors._TYPE_TO_EXTRACTOR)
    assert not missing, f"accepted but not extractable: {missing}"
    print("PASS  every accepted extension has an extractor "
          f"({len(accepted)} checked)")


def test_xlsm_is_routed_and_read():
    from openpyxl import Workbook

    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "book.xlsm")
        book = Workbook()
        book.active["A1"] = "needle in macro book"
        book.save(path)  # openpyxl writes the Open XML layout

        assert extractors.chunks_for_path(path) == \
            [("cell", "Sheet!A1", "needle in macro book")]
    print("PASS  .xlsm is routed to the Open XML extractor")


def _scan_with(folder, query):
    rows = []
    worker = ScanWorker([os.path.normpath(folder)], query, "exact", 85,
                        dict(TYPES), I18n("en"))
    worker.result.connect(rows.append)
    worker.run()
    return rows


def test_end_to_end_search_across_spreadsheet_formats():
    from openpyxl import Workbook

    with tempfile.TemporaryDirectory() as tmp:
        book = Workbook()
        book.active["A1"] = "findme in xlsx"
        book.save(os.path.join(tmp, "plain.xlsx"))

        book = Workbook()
        book.active["B2"] = "findme in xlsm"
        book.save(os.path.join(tmp, "macro.xlsm"))

        with open(os.path.join(tmp, "notes.txt"), "w",
                  encoding="utf-8") as handle:
            handle.write("findme in txt\n")

        rows = {row["file"]: row for row in _scan_with(tmp, "findme")}
        assert set(rows) == {"plain.xlsx", "macro.xlsm", "notes.txt"}, rows
        assert rows["macro.xlsm"]["locations"] == "Sheet!B2", \
            rows["macro.xlsm"]
        assert rows["plain.xlsx"]["locations"] == "Sheet!A1", \
            rows["plain.xlsx"]
    print("PASS  end-to-end scan finds text in .xlsx, .xlsm and .txt")


def test_legacy_xls_search():
    try:
        import xlwt
    except ImportError:
        print("SKIP  legacy .xls check (xlwt not installed)")
        return

    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "legacy.xls")
        workbook = xlwt.Workbook()
        sheet = workbook.add_sheet("Sayfa1")
        sheet.write(0, 0, "findme in xls")
        workbook.save(path)

        rows = _scan_with(tmp, "findme")
        assert len(rows) == 1, rows
        assert rows[0]["file"] == "legacy.xls", rows[0]
        assert rows[0]["locations"] == "Sayfa1!R1C1", rows[0]
    print("PASS  end-to-end scan finds text in .xls")


def test_unknown_extension_is_ignored():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "binary.bin")
        with open(path, "wb") as handle:
            handle.write(b"findme")
        assert extractors.chunks_for_path(path) is None
        assert _scan_with(tmp, "findme") == []
    print("PASS  unsupported extensions are still skipped")


if __name__ == "__main__":
    test_every_accepted_extension_has_an_extractor()
    test_xlsm_is_routed_and_read()
    test_end_to_end_search_across_spreadsheet_formats()
    test_legacy_xls_search()
    test_unknown_extension_is_ignored()
    print("ALL CHECKS PASSED")
